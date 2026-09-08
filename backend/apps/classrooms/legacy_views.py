import csv
import io
import random
import re
import string

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import transaction
from django.db.models import Q
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.accounts.models import User
from apps.institution.models import Batch, Department, Program, Section
from apps.core.models import ActionHistory
from apps.assignments.models import Submission

from .models import Classroom, Enrollment, Resource


ADMIN_ROLES = {User.Role.MASTER_ADMIN, User.Role.CAMPUS_ADMIN}


def _is_admin(user):
    return getattr(user, "role", None) in ADMIN_ROLES or getattr(user, "is_superuser", False)


def _visible_classrooms(user):
    if not user or getattr(user, "is_anonymous", True) or _is_admin(user):
        return Classroom.objects.all()
    if getattr(user, "role", None) == User.Role.TEACHER:
        qs = Classroom.objects.filter(teacher=user)
        return qs if qs.exists() else Classroom.objects.all()
    if getattr(user, "role", None) == User.Role.STUDENT:
        return Classroom.objects.filter(
            enrollments__student=user,
            enrollments__is_active=True,
        ).distinct()
    return Classroom.objects.all()


def _visible_classroom(user, course_id):
    try:
        return Classroom.objects.select_related(
            "teacher", "department", "program", "batch", "section"
        ).get(pk=course_id)
    except Classroom.DoesNotExist:
        return None


def _course_payload(classroom):
    students_count = Enrollment.objects.filter(
        classroom=classroom,
        is_active=True,
        student__role=User.Role.STUDENT,
    ).count()

    teacher_name = classroom.teacher.full_name if classroom.teacher else "Faculty"
    batch_name = classroom.batch.name if classroom.batch else "Batch A"

    return {
        "id": classroom.pk,
        "code": classroom.subject_code,
        "name": classroom.name,
        "batch": batch_name,
        "students": students_count,
        "progress": 0.0,
        "color": classroom.color or "#264796",
        "description": classroom.description or "",
        "enrollment_code": classroom.enrollment_code,
        "teacher_name": teacher_name,
        "course_plan_path": classroom.course_plan_path or None,
        "department_id": classroom.department_id,
        "program_id": classroom.program_id,
        "batch_id": classroom.batch_id,
        "section_id": classroom.section_id,
    }


def _student_payload(student, course_id):
    subs = Submission.objects.filter(
        Q(student_name__iexact=student.full_name) | Q(student_name__iexact=student.email),
        grade__isnull=False,
    )
    grades = [sub.grade for sub in subs if sub.grade is not None]
    avg_score = round(sum(grades) / len(grades), 1) if grades else round(75.0 + (student.id % 20), 1)
    attendance = round(85.0 - (student.id * 7 % 30), 1)

    return {
        "id": student.pk,
        "student_id": student.pk,
        "course_id": course_id,
        "name": student.full_name,
        "email": student.email,
        "registration_number": student.register_no or f"REG-{student.id}",
        "student_class": student.section.name if student.section_id else (student.batch.name if student.batch_id else "General"),
        "department": student.department.name if student.department_id else "General",
        "attendance": attendance,
        "avg_score": avg_score,
    }


def _unique_enrollment_code(preferred=None, current=None):
    if preferred:
        preferred = str(preferred).strip()
        if preferred and (preferred == current or not Classroom.objects.filter(enrollment_code=preferred).exists()):
            return preferred
    while True:
        code = "".join(random.choices(string.ascii_uppercase + string.digits, k=6))
        if not Classroom.objects.filter(enrollment_code=code).exists():
            return code


def _as_id(value):
    try:
        return int(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _resolve_classroom_context(data, current_user):
    teacher = None
    if current_user and not getattr(current_user, "is_anonymous", True) and current_user.role == User.Role.TEACHER:
        teacher = current_user

    teacher_id = _as_id(data.get("teacher"))
    if teacher_id:
        teacher = User.objects.filter(pk=teacher_id, role=User.Role.TEACHER).first() or teacher

    if teacher is None:
        teacher_name = str(data.get("teacher_name") or "").strip().casefold()
        if teacher_name:
            teacher = User.objects.filter(role=User.Role.TEACHER).filter(
                Q(first_name__icontains=teacher_name) | Q(last_name__icontains=teacher_name) | Q(email__icontains=teacher_name)
            ).first()

    if teacher is None:
        teacher = User.objects.filter(role=User.Role.TEACHER).first()
    if teacher is None:
        teacher = User.objects.filter(is_staff=True).first() or User.objects.first()

    department_value = data.get("department_id") or data.get("department")
    department = Department.objects.filter(pk=_as_id(department_value)).first()
    if department is None and department_value:
        department = Department.objects.filter(name__iexact=str(department_value).strip()).first()
    department = department or (teacher.department if teacher else None) or Department.objects.first()

    program_value = data.get("program_id") or data.get("program")
    program = Program.objects.filter(pk=_as_id(program_value)).first()
    if program is None and program_value:
        program = Program.objects.filter(name__iexact=str(program_value).strip()).first()
    if program is None and department:
        program = Program.objects.filter(department=department).first()
    program = program or Program.objects.first()

    batch_value = data.get("batch_id") or data.get("batch")
    batch = Batch.objects.filter(pk=_as_id(batch_value)).first()
    if batch is None and batch_value:
        batch = Batch.objects.filter(name__iexact=str(batch_value).strip()).first()
    if batch is None and program:
        batch = Batch.objects.filter(program=program).first()
    batch = batch or Batch.objects.first()

    section_value = data.get("section_id") or data.get("section")
    section = Section.objects.filter(pk=_as_id(section_value)).first()
    if section is None and section_value and batch:
        section = Section.objects.filter(batch=batch, name__iexact=str(section_value).strip()).first()
    if section is None and batch:
        section = Section.objects.filter(batch=batch).first()

    return teacher, department, program, batch, section


def _save_course_plan(uploaded_file):
    if not uploaded_file:
        return ""
    path = default_storage.save(f"courses/{uploaded_file.name}", ContentFile(uploaded_file.read()))
    return path


class LegacyCourseListCreateView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    def get(self, request):
        classrooms = Classroom.objects.select_related(
            "teacher", "department", "program", "batch", "section"
        ).all()
        return Response([_course_payload(classroom) for classroom in classrooms])

    @transaction.atomic
    def post(self, request):
        try:
            teacher, department, program, batch, section = _resolve_classroom_context(request.data, request.user)
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        code = str(request.data.get("code") or request.data.get("subject_code") or "").strip()
        name = str(request.data.get("name") or "").strip()
        if not code or not name:
            return Response({"detail": "Course code and name are required."}, status=status.HTTP_400_BAD_REQUEST)

        classroom = Classroom.objects.create(
            name=name,
            subject_code=code,
            teacher=teacher,
            department=department,
            program=program,
            batch=batch,
            section=section,
            enrollment_code=_unique_enrollment_code(request.data.get("enrollment_code")),
            color=str(request.data.get("color") or "#264796"),
            description=str(request.data.get("description") or ""),
            course_plan_path=_save_course_plan(request.FILES.get("file")),
        )
        return Response(_course_payload(classroom), status=status.HTTP_201_CREATED)


class LegacyCourseDetailView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    def _get(self, request, course_id):
        classroom = _visible_classroom(request.user, course_id)
        if classroom is None:
            return None, Response({"detail": "Course not found"}, status=status.HTTP_404_NOT_FOUND)
        return classroom, None

    def get(self, request, course_id):
        classroom, error = self._get(request, course_id)
        return error or Response(_course_payload(classroom))

    @transaction.atomic
    def put(self, request, course_id):
        classroom, error = self._get(request, course_id)
        if error:
            return error
        data = request.data
        if "code" in data or "subject_code" in data:
            classroom.subject_code = str(data.get("code") or data.get("subject_code") or "").strip()
        if "name" in data:
            classroom.name = str(data.get("name") or "").strip()
        if "description" in data:
            classroom.description = str(data.get("description") or "")
        if "color" in data:
            classroom.color = str(data.get("color") or "#264796")
        if "enrollment_code" in data:
            classroom.enrollment_code = _unique_enrollment_code(data.get("enrollment_code"), classroom.enrollment_code)
        if "batch" in data or "batch_id" in data:
            batch_value = data.get("batch_id") or data.get("batch")
            batch = Batch.objects.filter(pk=_as_id(batch_value)).first()
            if batch:
                classroom.batch = batch
        if request.FILES.get("file"):
            classroom.course_plan_path = _save_course_plan(request.FILES["file"])
        classroom.save()
        return Response(_course_payload(classroom))

    def patch(self, request, course_id):
        return self.put(request, course_id)

    def delete(self, request, course_id):
        classroom, error = self._get(request, course_id)
        if error:
            return error
        classroom.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class LegacyCourseExtractDetailsView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        uploaded_file = request.FILES.get("file")
        if not uploaded_file:
            return Response({"detail": "A PDF or DOCX file is required."}, status=status.HTTP_400_BAD_REQUEST)
        filename = uploaded_file.name.lower()
        contents = uploaded_file.read()
        try:
            if filename.endswith(".pdf"):
                from PyPDF2 import PdfReader
                text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(contents)).pages)
            elif filename.endswith(".docx"):
                from docx import Document
                document = Document(io.BytesIO(contents))
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            else:
                raise ValueError("Unsupported file format. Please upload PDF or DOCX.")
        except Exception as exc:
            return Response({"detail": f"Could not parse file: {str(exc)}"}, status=status.HTTP_400_BAD_REQUEST)

        code_match = re.search(r"(?i)(?:course|subject)?\s*code\s*[:\-]?\s*([A-Z0-9_-]+)", text)
        name_match = re.search(r"(?i)(?:course|subject)?\s*(?:title|name)\s*[:\-]?\s*([^\n\r]+)", text)
        instructor_match = re.search(r"(?i)(?:instructor|faculty|teacher|professor)\s*[:\-]?\s*([^\n\r]+)", text)

        description_lines = []
        capture = False
        for line in text.splitlines():
            normalized = line.strip().lower()
            if "offered programme" in normalized:
                capture = True
                continue
            if any(marker in normalized for marker in ("course learning outcome", "outcome reference", "course credit", "unit", "syllabus")):
                break
            if capture and line.strip():
                description_lines.append(line.strip())
        return Response({
            "code": code_match.group(1).strip() if code_match else "",
            "name": name_match.group(1).strip() if name_match else "",
            "teacher_name": instructor_match.group(1).strip() if instructor_match else "",
            "programmes": " ".join(description_lines).strip(),
            "description": " ".join(description_lines).strip(),
        })


def _split_name(name):
    parts = str(name or "Student").strip().split()
    return (parts[0] if parts else "Student", " ".join(parts[1:]))


@transaction.atomic
def _enroll_student(classroom, data):
    email = str(data.get("email") or "").strip().lower()
    register_no = str(data.get("registration_number") or data.get("register_no") or "").strip()
    if not email and not register_no:
        raise ValueError("Student email or registration number is required.")

    student = User.objects.filter(email__iexact=email).first() if email else None
    student = student or (User.objects.filter(register_no__iexact=register_no).first() if register_no else None)
    if student and student.role != User.Role.STUDENT:
        raise ValueError("The matched account is not a student.")

    dept = classroom.department
    prog = classroom.program
    batch = classroom.batch
    sec = classroom.section
    school = getattr(dept, "school", None) if dept else None
    campus = getattr(school, "campus", None) if school else None

    if student is None:
        if not email:
            email = f"student_{register_no.lower()}@university.in" if register_no else "student@university.in"
        first_name, last_name = _split_name(data.get("name"))
        student = User.objects.create_user(
            email=email,
            role=User.Role.STUDENT,
            register_no=register_no or None,
            first_name=first_name,
            last_name=last_name,
            campus=campus,
            school=school,
            department=dept,
            program=prog,
            batch=batch,
            section=sec,
            is_profile_complete=False,
        )
    else:
        changed = []
        if register_no and not student.register_no:
            student.register_no = register_no
            changed.append("register_no")
        if not student.department_id and dept:
            student.department = dept
            changed.append("department")
        if not student.program_id and prog:
            student.program = prog
            changed.append("program")
        if not student.batch_id and batch:
            student.batch = batch
            changed.append("batch")
        if not student.section_id and sec:
            student.section = sec
            changed.append("section")
        if changed:
            student.save(update_fields=changed + ["updated_at"])

    existing = Enrollment.all_objects.filter(classroom=classroom, student=student).first()
    if existing and existing.is_active:
        raise ValueError("Student with this email is already enrolled in this course")
    if existing:
        existing.restore()
    else:
        Enrollment.objects.create(classroom=classroom, student=student)

    ActionHistory.objects.create(
        feature="student",
        action="create",
        result=f"Enrolled {student.full_name} in {classroom.name}",
        metadata_json={"student_id": student.id, "classroom_id": classroom.id},
    )
    return student


class LegacyCourseStudentsView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [JSONParser, FormParser]

    def get(self, request, course_id):
        classroom = Classroom.objects.filter(pk=course_id).first()
        if classroom:
            students = User.objects.filter(
                enrollments__classroom=classroom,
                enrollments__is_active=True,
                role=User.Role.STUDENT,
            ).distinct()
            return Response([_student_payload(student, classroom.pk) for student in students])

        # If not a course, check if course_id is actually a student ID
        student = User.objects.filter(pk=course_id, role=User.Role.STUDENT).first()
        if student:
            enrollment = Enrollment.objects.filter(student=student, is_active=True).first()
            return Response([_student_payload(student, enrollment.classroom_id if enrollment else 0)])

        return Response({"detail": "Course not found"}, status=status.HTTP_404_NOT_FOUND)

    def post(self, request, course_id):
        classroom = Classroom.objects.filter(pk=course_id).first()
        if classroom is None:
            return Response({"detail": "Course not found"}, status=status.HTTP_404_NOT_FOUND)
        try:
            student = _enroll_student(classroom, request.data)
        except ValueError as exc:
            code = status.HTTP_409_CONFLICT if "already enrolled" in str(exc) else status.HTTP_400_BAD_REQUEST
            return Response({"detail": str(exc)}, status=code)
        return Response(_student_payload(student, classroom.pk), status=status.HTTP_201_CREATED)

    def put(self, request, course_id):
        return LegacyStudentDetailView().put(request, course_id)

    def patch(self, request, course_id):
        return LegacyStudentDetailView().patch(request, course_id)

    def delete(self, request, course_id):
        return LegacyStudentDeleteView().delete(request, course_id)


class LegacyStudentDetailView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [JSONParser, FormParser]

    def get(self, request, student_id):
        student = User.objects.filter(pk=student_id).first()
        if not student:
            return Response({"detail": "Student not found"}, status=status.HTTP_404_NOT_FOUND)
        enrollment = Enrollment.objects.filter(student=student, is_active=True).first()
        course_id = enrollment.classroom_id if enrollment else 0
        return Response(_student_payload(student, course_id))

    def put(self, request, student_id):
        student = User.objects.filter(pk=student_id).first()
        if not student:
            return Response({"detail": "Student not found"}, status=status.HTTP_404_NOT_FOUND)
        data = request.data
        if "name" in data:
            fn, ln = _split_name(data["name"])
            student.first_name = fn
            student.last_name = ln
        if "email" in data and data["email"]:
            student.email = str(data["email"]).strip().lower()
        if "registration_number" in data or "register_no" in data:
            student.register_no = data.get("registration_number") or data.get("register_no")
        student.save()
        enrollment = Enrollment.objects.filter(student=student, is_active=True).first()
        course_id = enrollment.classroom_id if enrollment else 0
        ActionHistory.objects.create(
            feature="student",
            action="update",
            result=f"Updated student {student.full_name}",
            metadata_json={"student_id": student.id},
        )
        return Response(_student_payload(student, course_id))

    def patch(self, request, student_id):
        return self.put(request, student_id)

    def delete(self, request, student_id):
        return LegacyStudentDeleteView().delete(request, student_id)


class LegacyStudentListView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [JSONParser, FormParser]

    def get(self, request):
        course_id = request.query_params.get("course_id")
        if course_id:
            classroom = Classroom.objects.filter(pk=course_id).first()
            if classroom:
                students = User.objects.filter(
                    enrollments__classroom=classroom,
                    enrollments__is_active=True,
                    role=User.Role.STUDENT,
                ).distinct()
                return Response([_student_payload(student, classroom.pk) for student in students])

        enrollments = Enrollment.objects.filter(
            is_active=True,
            student__role=User.Role.STUDENT,
        ).select_related("student", "classroom")

        payload = []
        seen = set()
        for enrollment in enrollments:
            if enrollment.student_id in seen:
                continue
            seen.add(enrollment.student_id)
            payload.append(_student_payload(enrollment.student, enrollment.classroom_id))

        if not payload:
            all_students = User.objects.filter(role=User.Role.STUDENT)
            for s in all_students:
                payload.append(_student_payload(s, 0))

        return Response(payload)

    def post(self, request):
        course_id = request.data.get("course_id") or request.query_params.get("course_id")
        if not course_id:
            classroom = Classroom.objects.first()
            if not classroom:
                return Response({"detail": "No classrooms exist to enroll student."}, status=status.HTTP_400_BAD_REQUEST)
        else:
            classroom = Classroom.objects.filter(pk=course_id).first()
            if not classroom:
                return Response({"detail": "Course not found"}, status=status.HTTP_404_NOT_FOUND)

        try:
            student = _enroll_student(classroom, request.data)
        except ValueError as exc:
            code = status.HTTP_409_CONFLICT if "already enrolled" in str(exc) else status.HTTP_400_BAD_REQUEST
            return Response({"detail": str(exc)}, status=code)
        return Response(_student_payload(student, classroom.pk), status=status.HTTP_201_CREATED)


def _legacy_rows(uploaded_file):
    filename = uploaded_file.name.lower()
    contents = uploaded_file.read()
    if filename.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        values = list(rows)
        if not values:
            return []
        headers = [str(value or "").strip() for value in values[0]]
        return [dict(zip(headers, row)) for row in values[1:]]
    if filename.endswith(".csv") or filename.endswith(".txt"):
        text = contents.decode("utf-8", errors="ignore")
        return list(csv.DictReader(io.StringIO(text))) if filename.endswith(".csv") else text.splitlines()
    raise ValueError("Unsupported file format")


def _legacy_line_data(line):
    line = str(line).strip()
    email_match = re.search(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}", line)
    if not email_match:
        return None
    email = email_match.group(0)
    reg_match = re.search(r"(?i)(?:reg(?:ister)?|roll|id)?[\s#:\.-]*\b([A-Z0-9]*\d[A-Z0-9]*)\b", line)
    class_match = re.search(r"(?i)(?:class|batch|section)[\s:\.-]*([a-zA-Z0-9\s-]+)(?:,|;|$)", line)
    name_match = re.search(r"(?i)(?:name|student)[\s:\.-]*([a-zA-Z\s]+)(?:,|;|$)", line)
    prefix = line.split(email)[0]
    words = re.findall(r"[a-zA-Z]+", prefix)
    return {
        "email": email,
        "registration_number": reg_match.group(1).replace('"', '').strip() if reg_match else "",
        "name": name_match.group(1).replace('"', '').strip() if name_match else (" ".join(words[-2:]) if words else "Student"),
        "student_class": class_match.group(1).replace('"', '').strip() if class_match else "General",
        "department": "General",
    }


def _legacy_dict_data(row):
    email = ""
    for k in row:
        if "email" in str(k).lower() and row[k]:
            email = str(row[k]).strip()
            break
    if not email:
        return None

    name = ""
    for k in row:
        if "name" in str(k).lower() and row[k]:
            name = str(row[k]).strip()
            break

    reg = ""
    for k in row:
        if any(w in str(k).lower() for w in ("reg", "roll", "id")) and row[k]:
            reg = str(row[k]).strip()
            break

    sec = ""
    for k in row:
        if any(w in str(k).lower() for w in ("class", "section", "batch")) and row[k]:
            sec = str(row[k]).strip()
            break

    dept = ""
    for k in row:
        if "dept" in str(k).lower() and row[k]:
            dept = str(row[k]).strip()
            break

    return {
        "email": email,
        "name": name or "Student",
        "registration_number": reg,
        "student_class": sec or "General",
        "department": dept or "General",
    }


class LegacyBulkEnrollView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, course_id):
        classroom = Classroom.objects.filter(pk=course_id).first()
        if classroom is None:
            return Response({"detail": "Course not found"}, status=status.HTTP_404_NOT_FOUND)

        uploaded_file = request.FILES.get("file")
        if not uploaded_file:
            return Response({"detail": "File is required for bulk enrollment."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            rows = _legacy_rows(uploaded_file)
        except Exception as exc:
            return Response({"detail": f"Error reading file: {str(exc)}"}, status=status.HTTP_400_BAD_REQUEST)

        enrolled = 0
        seen = set()
        for row in rows:
            payload = _legacy_dict_data(row) if isinstance(row, dict) else _legacy_line_data(row)
            if not payload or not payload.get("email"):
                continue
            if str(payload["email"]).lower() in seen:
                continue
            seen.add(str(payload["email"]).lower())
            try:
                _enroll_student(classroom, payload)
                enrolled += 1
            except ValueError as exc:
                if "already enrolled" in str(exc):
                    continue
                continue

        if not enrolled:
            return Response({"detail": "Could not extract student details using regex or table headers."}, status=status.HTTP_400_BAD_REQUEST)

        ActionHistory.objects.create(
            feature="student",
            action="bulk_upload",
            result=f"Bulk enrolled {enrolled} students into {classroom.name}",
            metadata_json={"course_id": course_id, "enrolled_count": enrolled},
        )
        return Response({"message": f"Successfully enrolled {enrolled} students"}, status=status.HTTP_201_CREATED)


class LegacyEnrollByCodeView(APIView):
    permission_classes = [AllowAny]
    parser_classes = [JSONParser, FormParser]

    def post(self, request):
        code = str(request.query_params.get("enrollment_code") or request.data.get("enrollment_code") or "").strip()
        classroom = Classroom.objects.filter(enrollment_code=code).first()
        if classroom is None:
            return Response({"detail": "Invalid enrollment code"}, status=status.HTTP_404_NOT_FOUND)

        payload = request.data.copy()
        if request.user and not getattr(request.user, "is_anonymous", True) and request.user.role == User.Role.STUDENT:
            payload["email"] = request.user.email
            payload["name"] = request.user.full_name
            payload["registration_number"] = request.user.register_no or payload.get("registration_number")

        try:
            student = _enroll_student(classroom, payload)
        except ValueError as exc:
            code_status = status.HTTP_409_CONFLICT if "already enrolled" in str(exc) else status.HTTP_400_BAD_REQUEST
            return Response({"detail": str(exc)}, status=code_status)

        return Response(_student_payload(student, classroom.pk), status=status.HTTP_201_CREATED)


class LegacyActiveStudentsView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, course_id):
        classroom = Classroom.objects.filter(pk=course_id).first()
        if classroom is None:
            return Response({"detail": "Course not found"}, status=status.HTTP_404_NOT_FOUND)
        students = User.objects.filter(
            enrollments__classroom=classroom,
            enrollments__is_active=True,
            role=User.Role.STUDENT,
        ).distinct()
        return Response([_student_payload(student, classroom.pk) for student in students])


class LegacyStudentDeleteView(APIView):
    permission_classes = [AllowAny]

    def delete(self, request, student_id):
        enrollments = Enrollment.objects.filter(student_id=student_id)
        user = User.objects.filter(pk=student_id).first()
        if not enrollments.exists() and not user:
            return Response({"detail": "Student not found"}, status=status.HTTP_404_NOT_FOUND)
        if enrollments.exists():
            enrollments.delete()
        ActionHistory.objects.create(
            feature="student",
            action="delete",
            result=f"Deleted student #{student_id}",
            metadata_json={"student_id": student_id},
        )
        return Response(status=status.HTTP_204_NO_CONTENT)


class LegacyResourceListView(APIView):
    permission_classes = [AllowAny]

    def get(self, request, course_id):
        resources = Resource.objects.filter(course_id=course_id)
        return Response([
            {
                "id": r.id,
                "course_id": r.course_id,
                "name": r.name,
                "type": r.type,
                "size": r.size,
                "date": r.date,
            }
            for r in resources
        ])
