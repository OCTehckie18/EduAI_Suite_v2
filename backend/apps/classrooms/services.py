import csv
import io
from dataclasses import dataclass

from django.db import transaction
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.institution.models import Section

from .models import Classroom, Enrollment


EXPECTED_COLUMNS = ("register no", "student name", "student email", "section")


@dataclass
class EnrollmentImportResult:
    enrolled: int = 0
    already_enrolled: int = 0
    provisioned: int = 0
    errors: list | None = None

    def as_dict(self):
        return {
            "enrolled": self.enrolled,
            "already_enrolled": self.already_enrolled,
            "provisioned": self.provisioned,
            "errors": self.errors or [],
        }


def _normalise(value):
    return " ".join(str(value or "").strip().lower().split())


def _rows_from_file(uploaded_file):
    filename = uploaded_file.name.lower()
    if filename.endswith(".csv"):
        text = uploaded_file.read().decode("utf-8-sig")
        rows = csv.reader(io.StringIO(text))
        headers = [_normalise(value) for value in next(rows, ())]
        return ({headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))} for values in rows)
    if filename.endswith(".xlsx"):
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [_normalise(value) for value in next(rows, ())]
        return ({headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))} for values in rows)
    raise ValueError("Only .csv and .xlsx files are supported.")


def _student_for_row(row, classroom):
    email = _normalise(row.get("student email"))
    register_no = str(row.get("register no") or "").strip()
    if not email and not register_no:
        raise ValueError("Student Email or Register No is required.")

    student = User.objects.filter(register_no__iexact=register_no).first() if register_no else None
    if student is None and email:
        student = User.objects.filter(email__iexact=email).first()
    if student:
        if student.role != User.Role.STUDENT:
            raise ValueError("The matched account is not a student.")
        return student, False

    if not email:
        raise ValueError("A student email is required for placeholder account creation.")
    student = User.objects.create_user(
        email=email,
        role=User.Role.STUDENT,
        register_no=register_no or None,
        first_name=str(row.get("student name") or "").strip(),
        campus=classroom.batch.program.department.school.campus,
        school=classroom.batch.program.department.school,
        department=classroom.department,
        program=classroom.program,
        batch=classroom.batch,
        is_profile_complete=False,
    )
    return student, True


def _section_for_row(row, classroom):
    section_name = str(row.get("section") or "").strip()
    if classroom.section:
        if section_name and section_name.casefold() != classroom.section.name.casefold():
            raise ValueError("The row section does not match this classroom section.")
        return classroom.section
    if section_name:
        try:
            return Section.objects.get(batch=classroom.batch, name__iexact=section_name)
        except Section.DoesNotExist as exc:
            raise ValueError(f"Section '{section_name}' does not belong to this batch.") from exc
    return None


@transaction.atomic
def import_enrollments(classroom: Classroom, uploaded_file):
    result = EnrollmentImportResult(errors=[])
    rows = _rows_from_file(uploaded_file)
    for row_number, row in enumerate(rows, start=2):
        try:
            student, provisioned = _student_for_row(row, classroom)
            section = _section_for_row(row, classroom)
            if section and student.section_id and student.section_id != section.id:
                raise ValueError("Student belongs to a different section.")
            if section and not student.section_id:
                student.section = section
                student.save(update_fields=["section", "updated_at"])
            enrollment = Enrollment.all_objects.filter(classroom=classroom, student=student).first()
            if enrollment and enrollment.is_active:
                result.already_enrolled += 1
                continue
            if section and not enrollment and Enrollment.objects.filter(classroom=classroom, student__section=section).count() >= section.max_capacity:
                raise ValueError(f"Section '{section.name}' has reached its maximum capacity.")
            if enrollment:
                enrollment.restore()
            else:
                Enrollment.objects.create(classroom=classroom, student=student)
            result.enrolled += 1
            result.provisioned += int(provisioned)
        except (ValueError, TypeError) as exc:
            result.errors.append({"row": row_number, "error": str(exc)})
    return result
